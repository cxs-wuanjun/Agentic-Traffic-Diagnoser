import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report(content: str, chart_path: str = None, sop_content: str = None) -> str:
    """
    统一的红头文件排版组装引擎
    返回生成的 Word 文档绝对路径
    """
    doc = Document()
    
    # 统一红头标题配置
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("交通智能诊断评估报告")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    doc.add_paragraph("").add_run("_" * 60)

    # 1. 正文分析插入
    doc.add_heading("1. 诊断分析", level=1)
    doc.add_paragraph(content)

    # 2. 动态插入图表
    if chart_path and os.path.exists(chart_path):
        doc.add_heading("2. 数据可视化", level=1)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(chart_path, width=Inches(6.0))

    # 3. 动态插入合规预案
    if sop_content:
        doc.add_heading("3. 官方处置建议", level=1)
        doc.add_paragraph(sop_content)

    out_dir = os.path.join(os.getcwd(), "outputs", "reports")
    os.makedirs(out_dir, exist_ok=True)
    out_file = os.path.join(out_dir, "Final_Report.docx")
    
    doc.save(out_file)
    return out_file
